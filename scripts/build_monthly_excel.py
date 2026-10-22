"""Build the governed, refreshable monthly management workbook.

Usage: python scripts/build_monthly_excel.py --month 2026-09
"""
from __future__ import annotations

import argparse
import calendar
from datetime import date, datetime
from pathlib import Path
import re
import os

import duckdb
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
DB = Path(os.environ.get("PULSEIQ_DB_PATH", ROOT / "pulseiq.duckdb"))
OUT = ROOT / "outputs" / "pulseiq-monthly"
NAVY = "14213D"
TEAL = "117C78"
PALE = "E8F3F1"
INK = "243247"
MUTED = "64748B"
WHITE = "FFFFFF"
RED_FILL = "FCE8E6"
RED_TEXT = "B42318"
AMBER_FILL = "FFF4D6"
AMBER_TEXT = "8A5200"
GREEN_FILL = "E7F5EC"
GREEN_TEXT = "18794E"
THIN_GREY = Side(style="thin", color="D9E2EC")


def query(con: duckdb.DuckDBPyConnection, sql: str, params: list | None = None) -> pd.DataFrame:
    return con.execute(sql, params or []).fetchdf()


def scalar(con: duckdb.DuckDBPyConnection, sql: str, params: list | None = None):
    rows = con.execute(sql, params or []).fetchone()
    return None if rows is None else rows[0]


def month_bounds(month: date) -> tuple[date, date]:
    first = month.replace(day=1)
    last = first.replace(day=calendar.monthrange(first.year, first.month)[1])
    return first, last


def _excel_value(value):
    if value is None or pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, (date, datetime, str, int, float, bool)):
        return value
    return str(value)


def _number_format(header: str, value) -> str | None:
    name = header.lower()
    if isinstance(value, (date, datetime, pd.Timestamp)):
        return "yyyy-mm-dd"
    if "period" in name or name.endswith("_date") or name in {"as_of_date", "source data cutoff"}:
        return "yyyy-mm-dd"
    if "reconciliation_difference" in name or name == "difference":
        return None
    if name in {"variance_pct", "operating_margin_pct", "completion_pct", "on_time_pct", "monthly attrition"}:
        return '0.00"%"'
    if "currency" in name or any(x in name for x in ("cost", "pay", "revenue", "income", "expense", "profit", "balance", "cash", "variance", "target_value", "actual_value", "difference")):
        return '#,##0.00;[Red](#,##0.00);-'
    if any(x in name for x in ("rows", "headcount", "hires", "exits", "invoices", "tasks", "count", "sample_size", "velocity", "bad_rows")):
        return '#,##0;[Red](#,##0);-'
    if "days" in name or "cycle_time" in name or name == "dso":
        return '#,##0.0;[Red](#,##0.0);-'
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return '#,##0.00;[Red](#,##0.00);-'
    return None


def _base_sheet(wb: Workbook, name: str, title: str, context: str):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = TEAL if name == "Overview" else NAVY
    ws["A1"] = title
    ws["A1"].font = Font(name="Arial", size=15, bold=True, color=NAVY)
    ws["A2"] = context
    ws["A2"].font = Font(name="Arial", size=9, italic=True, color=MUTED)
    ws.freeze_panes = "A5"
    ws.row_dimensions[1].height = 23
    ws.row_dimensions[2].height = 20
    return ws


def _write_table(ws, start_row: int, headers: list[str], frame: pd.DataFrame, table_name: str) -> int:
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(start_row, col_idx, header)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Arial", size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin", color=WHITE))
    ws.row_dimensions[start_row].height = 36 if any(" " in h or len(h) > 14 for h in headers) else 30

    for row_idx, row in enumerate(frame.itertuples(index=False, name=None), start_row + 1):
        row_height = 22
        for col_idx, (header, raw) in enumerate(zip(headers, row), 1):
            value = _excel_value(raw)
            cell = ws.cell(row_idx, col_idx, value)
            cell.font = Font(name="Arial", size=9, color=INK)
            cell.alignment = Alignment(
                horizontal="left" if isinstance(value, str) else "right",
                vertical="center",
                wrap_text=header.lower() in {"evidence / note", "recommendation", "check_name", "table_name"},
            )
            fmt = _number_format(header, value)
            if fmt:
                cell.number_format = fmt
            if header.lower() in {"evidence / note", "recommendation", "check_name", "table_name"}:
                row_height = max(row_height, 40 if header.lower() == "check_name" else 50 if header.lower() == "recommendation" else 34)
            if row_idx % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F5F8FA")
            cell.border = Border(bottom=THIN_GREY)
        ws.row_dimensions[row_idx].height = row_height

    last_row = start_row + len(frame)
    last_col = len(headers)
    if len(frame):
        ref = f"A{start_row}:{ws.cell(start_row, last_col).column_letter}{last_row}"
        table = Table(displayName=table_name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False,
            showLastColumn=False, showRowStripes=True, showColumnStripes=False,
        )
        ws.add_table(table)
    for col_idx, header in enumerate(headers, 1):
        values = [str(header)] + [str(_excel_value(v) or "") for v in frame.iloc[:, col_idx - 1].tolist()]
        width = max(12, min(44, max((len(v) for v in values), default=12) + 2))
        if header.lower() in {"evidence / note", "recommendation", "check_name", "table_name"}:
            width = min(52, max(width, 26))
        ws.column_dimensions[ws.cell(1, col_idx).column_letter].width = width
    return last_row


def _style_status(ws, column: int, first_row: int, last_row: int):
    if last_row < first_row:
        return
    letter = ws.cell(1, column).column_letter
    target = f"{letter}{first_row}:{letter}{last_row}"
    for text, fill, color in (
        ("FAIL", RED_FILL, RED_TEXT),
        ("REVIEW REQUIRED", RED_FILL, RED_TEXT),
        ("WARN", AMBER_FILL, AMBER_TEXT),
        ("AMBER", AMBER_FILL, AMBER_TEXT),
        ("PASS", GREEN_FILL, GREEN_TEXT),
        ("GREEN", GREEN_FILL, GREEN_TEXT),
    ):
        ws.conditional_formatting.add(
            target,
            CellIsRule(operator="equal", formula=[f'"{text}"'],
                       fill=PatternFill("solid", fgColor=fill),
                       font=Font(name="Arial", size=9, bold=True, color=color)),
        )


def _metadata(month: date, cutoff: date, source: str) -> str:
    return f"Report month: {month:%Y-%m}  |  Data cutoff: {cutoff:%Y-%m-%d}  |  Source: {source}"


def build(month: date, output: Path, db_path: Path = DB) -> Path:
    if not db_path.exists():
        raise FileNotFoundError(f"Warehouse missing at {db_path}; run python etl.py first.")
    month_start, month_end = month_bounds(month)
    con = duckdb.connect(str(db_path), read_only=True)
    try:
        exists = scalar(con, "SELECT COUNT(*) FROM metrics.m01_headcount WHERE period=? AND department_id='ALL'", [month_start])
        if not exists:
            raise ValueError(f"No workforce data for {month_start:%Y-%m}; choose a month in the observed range.")
        cutoff = scalar(con, "SELECT data_cutoff FROM metrics.dashboard_data_freshness")
        latest_run = scalar(con, "SELECT run_id FROM audit.etl_run_log ORDER BY started_at DESC LIMIT 1")

        workforce = scalar(con, "SELECT value FROM metrics.m01_headcount WHERE period=? AND department_id='ALL'", [month_start])
        attrition = scalar(con, "SELECT value FROM metrics.m04_attrition_monthly WHERE period=? AND department_id='ALL'", [month_start])
        employer_cost = scalar(con, "SELECT value FROM metrics.m09_payroll_cost WHERE period=? AND department_id='ALL'", [month_start])
        net_profit = scalar(con, "SELECT net_profit FROM metrics.dashboard_pnl_monthly WHERE period=?", [month_start])
        ar_total = scalar(con, "SELECT SUM(balance) FROM metrics.dashboard_receivables_asof WHERE as_of_date=?", [month_end])
        dso = scalar(con, "SELECT value FROM metrics.m20_dso WHERE period=?", [month_end])
        cash = query(con, "SELECT cash_in,cash_out,net_movement,closing_balance FROM metrics.dashboard_cash_movement WHERE period=?", [month_start])
        cash_row = cash.iloc[0].tolist() if not cash.empty else [None, None, None, None]
        run_status = scalar(con, "SELECT status FROM audit.etl_run_log WHERE run_id=?", [latest_run]) if latest_run else None

        budget = query(con, """SELECT department_id,budget_cost,actual_cost,variance,variance_pct,status,recommendation
            FROM metrics.dashboard_budget_actuals WHERE period=? ORDER BY variance DESC""", [month_start])
        kpis = query(con, """SELECT department_id,metric,target_value,actual_value,direction,sample_size,status,variance
            FROM metrics.dashboard_kpi_targets WHERE period=? ORDER BY department_id,metric""", [month_start])
        payroll = query(con, """SELECT p.department_id,g.value AS gross_pay,n.value AS net_pay,
            p.value AS employer_cost,h.value AS cost_per_head
            FROM metrics.m09_payroll_cost p
            JOIN metrics.m07_gross_pay g USING(period,department_id)
            JOIN metrics.m08_net_pay n USING(period,department_id)
            JOIN metrics.m10_cost_per_head h USING(period,department_id)
            WHERE p.period=? AND p.department_id<>'ALL' ORDER BY p.value DESC""", [month_start])
        pnl = query(con, """SELECT period,revenue,other_income,expenses,net_profit,operating_margin_pct,reconciliation_difference
            FROM metrics.dashboard_pnl_monthly ORDER BY period""")
        receivables = query(con, """SELECT ageing_bucket,COUNT(DISTINCT invoice_id) AS invoices,SUM(balance) AS balance
            FROM metrics.dashboard_receivables_asof WHERE as_of_date=? GROUP BY 1
            ORDER BY CASE ageing_bucket WHEN 'Not due' THEN 0 WHEN '1-30 days' THEN 1
                WHEN '31-60 days' THEN 2 WHEN '61-90 days' THEN 3 ELSE 4 END""", [month_end])
        cash_history = query(con, """SELECT period,cash_in,cash_out,net_movement,closing_balance
            FROM metrics.dashboard_cash_movement ORDER BY period""")
        teams = query(con, """SELECT sprint_id,team_id,team_name,department_id,planned_tasks,done_tasks,
            overdue_tasks,blocked_tasks,reopened_count,completion_pct,on_time_pct,velocity,cycle_time_days
            FROM metrics.dashboard_task_operations ORDER BY sprint_id,team_id""")
        dq = query(con, """SELECT check_id,check_name,severity,bad_rows,status
            FROM audit.dq_results WHERE run_id=? ORDER BY check_id""", [latest_run])
        load_summary = query(con, """SELECT table_name,rows_in,rows_passed,rows_quarantined
            FROM audit.load_summary WHERE run_id=? ORDER BY table_name""", [latest_run])
        reconciliation = query(con, """SELECT period,payroll_cost,ledger_payroll_cost,difference,status
            FROM audit.payroll_ledger_reconciliation_clean ORDER BY period""")
    finally:
        con.close()

    wb = Workbook()
    wb.remove(wb.active)
    meta = _metadata(month_start, cutoff, "governed metric, dashboard and audit views")

    ws = _base_sheet(wb, "Overview", "PulseIQ monthly management pack", meta)
    overview = pd.DataFrame([
        ["Reporting month", month_start, "month", "Latest observed workforce period"],
        ["Source data cutoff", cutoff, "date", "Latest supplied source date"],
        ["Month-end headcount", workforce, "people", "Governed metric M01"],
        ["Monthly attrition", attrition, "percent points", "Governed metric M04; exits divided by starting headcount"],
        ["Employer payroll cost", employer_cost, "INR", "Governed metric M09"],
        ["Net profit", net_profit, "INR", "Ledger P&L"],
        ["Open receivables", ar_total, "INR", "Positive open invoices as of month-end"],
        ["Days sales outstanding", dso, "days", "Governed metric M20"],
        ["Cash in", cash_row[0], "INR", "Ledger movements; no opening bank balance supplied"],
        ["Cash out", cash_row[1], "INR", "Ledger movements; no opening bank balance supplied"],
        ["Net cash movement", cash_row[2], "INR", "Ledger movement, not an ending bank balance"],
        ["ETL run", latest_run, "run id", f"Latest warehouse refresh; status {run_status or 'unknown'}"],
        ["Data quality gate", "REVIEW REQUIRED" if not dq.empty and (dq.status.isin(["FAIL", "WARN"]).any()) else "PASS",
         "status", "Review seeded raw-source findings before external release"],
    ], columns=["Measure", "Value", "Unit", "Evidence / note"])
    _write_table(ws, 4, overview.columns.tolist(), overview, "OverviewMetrics")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 64
    for row in range(5, 5 + len(overview)):
        unit = ws.cell(row, 3).value
        if unit == "INR":
            ws.cell(row, 2).number_format = '#,##0.00;[Red](#,##0.00);-'
        elif unit == "percent points":
            ws.cell(row, 2).number_format = '0.00"%"'
        elif unit == "days":
            ws.cell(row, 2).number_format = "0.0"
        elif unit == "people":
            ws.cell(row, 2).number_format = "#,##0"
    _style_status(ws, 2, 17, 17)

    ws = _base_sheet(wb, "Budget and KPI", "Payroll budget and KPI target status", meta)
    last = _write_table(ws, 4, budget.columns.tolist(), budget, "DepartmentBudget")
    kpi_title_row = last + 3
    ws.cell(kpi_title_row, 1, "KPI target status")
    ws.cell(kpi_title_row, 1).font = Font(name="Arial", size=12, bold=True, color=TEAL)
    kpi_header_row = kpi_title_row + 1
    _write_table(ws, kpi_header_row, kpis.columns.tolist(), kpis, "KpiTargets")
    for row in range(kpi_header_row + 1, kpi_header_row + len(kpis) + 1):
        if str(ws.cell(row, 2).value).lower() in {"attendance rate", "monthly attrition", "overtime share", "task completion"}:
            for col in (3, 4, 8):
                ws.cell(row, col).number_format = '0.00"%";[Red](0.00"%");-'
    _style_status(ws, 6, 5, last)
    ws.column_dimensions["G"].width = 80
    for row in range(4, last + 1):
        ws.cell(row, 7).alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
    _style_status(ws, 7, kpi_header_row + 1, kpi_header_row + len(kpis))

    ws = _base_sheet(wb, "Payroll", "Payroll by department", meta)
    _write_table(ws, 4, payroll.columns.tolist(), payroll, "PayrollByDepartment")

    ws = _base_sheet(wb, "P and L", "Monthly profit and loss", meta)
    _write_table(ws, 4, pnl.columns.tolist(), pnl, "MonthlyPnL")

    ws = _base_sheet(wb, "Receivables", "Open receivables by ageing bucket", meta)
    _write_table(ws, 4, receivables.columns.tolist(), receivables, "ReceivablesAgeing")

    ws = _base_sheet(wb, "Cash", "Cash movements by month", meta + "; ledger movements, not an opening-bank-balance roll-forward")
    _write_table(ws, 4, cash_history.columns.tolist(), cash_history, "CashMovements")

    ws = _base_sheet(wb, "Team Operations", "Team productivity across all sprints", meta + "; task metrics use the clean task feed")
    _write_table(ws, 4, teams.columns.tolist(), teams, "TeamOperations")

    ws = _base_sheet(wb, "Quality", "Source quality findings and row reconciliation", meta + "; review seeded source findings before any external release")
    last = _write_table(ws, 4, dq.columns.tolist(), dq, "QualityFindings")
    _style_status(ws, 5, 5, last)
    load_title_row = last + 3
    ws.cell(load_title_row, 1, "Source row reconciliation")
    ws.cell(load_title_row, 1).font = Font(name="Arial", size=12, bold=True, color=TEAL)
    load_header_row = load_title_row + 1
    load_last = _write_table(ws, load_header_row, load_summary.columns.tolist(), load_summary, "SourceReconciliation")
    if load_last >= load_header_row + 1:
        first, final = load_header_row + 1, load_last
        letter = ws.cell(1, 4).column_letter
        ws.conditional_formatting.add(
            f"{letter}{first}:{letter}{final}",
            CellIsRule(operator="greaterThan", formula=["0"],
                       fill=PatternFill("solid", fgColor=AMBER_FILL),
                       font=Font(name="Arial", size=9, color=AMBER_TEXT)),
        )

    ws = _base_sheet(wb, "Reconciliation", "Clean payroll to ledger reconciliation", meta + "; 18 monthly periods")
    last = _write_table(ws, 4, reconciliation.columns.tolist(), reconciliation, "PayrollLedgerReconciliation")
    _style_status(ws, 5, 5, last)

    for ws in wb.worksheets:
        ws.sheet_view.zoomScale = 90
        ws.sheet_properties.pageSetUpPr.fitToPage = False
        ws.page_setup.fitToHeight = 0
        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.sheet_view.showGridLines = False

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output


def main():
    parser = argparse.ArgumentParser(description="Generate a monthly PulseIQ Excel pack from governed views.")
    parser.add_argument("--month", help="Report month as YYYY-MM; defaults to the latest observed workforce period.")
    parser.add_argument("--output", type=Path, help="Optional output workbook path.")
    parser.add_argument("--db", type=Path, default=DB, help="DuckDB warehouse path.")
    args = parser.parse_args()
    if not args.db.exists():
        raise SystemExit(f"Warehouse missing at {args.db}; run python etl.py first.")
    if args.month:
        if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", args.month):
            raise SystemExit("--month must be YYYY-MM.")
        year, month_num = map(int, args.month.split("-"))
        month = date(year, month_num, 1)
    else:
        with duckdb.connect(str(args.db), read_only=True) as con:
            month = scalar(con, "SELECT MAX(period) FROM metrics.m01_headcount WHERE department_id='ALL'")
    output = args.output or OUT / f"PulseIQ_Monthly_Pack_{month:%Y-%m}.xlsx"
    path = build(month, output, args.db)
    print(f"Workbook written to {path}")


if __name__ == "__main__":
    main()
