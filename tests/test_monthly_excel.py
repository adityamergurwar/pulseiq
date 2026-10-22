import tempfile
import unittest
from datetime import date
from pathlib import Path

import duckdb
from openpyxl import load_workbook

from scripts.build_monthly_excel import DB, ROOT, build


class MonthlyExcelTests(unittest.TestCase):
    def test_monthly_export_has_expected_views_and_governed_values(self):
        if not DB.exists():
            self.skipTest("Run python etl.py to build the local warehouse first.")
        con = duckdb.connect(str(DB), read_only=True)
        try:
            expected_headcount = con.execute(
                "SELECT value FROM metrics.m01_headcount WHERE period='2026-09-01' AND department_id='ALL'"
            ).fetchone()[0]
        finally:
            con.close()
        build_tmp = ROOT / ".artifact_build"
        build_tmp.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=build_tmp) as directory:
            output = Path(directory) / "PulseIQ_Monthly_Pack_2026-09.xlsx"
            build(date(2026, 9, 1), output, DB)
            workbook = load_workbook(output, data_only=True)
            self.assertEqual(
                workbook.sheetnames,
                ["Overview", "Budget and KPI", "Payroll", "P and L", "Receivables", "Cash", "Team Operations", "Quality", "Reconciliation"],
            )
            self.assertEqual(workbook["Overview"]["B7"].value, expected_headcount)
            self.assertEqual(workbook["Overview"]["B17"].value, "REVIEW REQUIRED")
            self.assertEqual(workbook["Quality"]["E5"].value, "FAIL")
            self.assertEqual(workbook["Reconciliation"]["E22"].value, "PASS")
            self.assertEqual(workbook["Cash"]["A2"].value.split("|")[0].strip(), "Report month: 2026-09")
            self.assertTrue(all(ws.freeze_panes == "A5" for ws in workbook.worksheets))
            self.assertTrue(all(ws.tables for ws in workbook.worksheets))
            workbook.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
